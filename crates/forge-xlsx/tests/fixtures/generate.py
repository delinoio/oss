"""External XLSX fixtures authored with openpyxl 3.1.5 (test-only)."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from datetime import datetime
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule, ColorScaleRule, DataBarRule, IconSetRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

path = Path(__file__).with_name("external.xlsx")
book = Workbook()
sheet = book.active
sheet.title = "External"
sheet.append(["Category", "Values", "Other"])
for row in range(2, 7):
    sheet.append([f"Item {row}", row * 2, row * 3])
sheet["E1"] = "External text"
sheet["E1"].font = Font(name="Arial", bold=True, color="112233")
sheet["E2"] = "=SUM(B2:B6)"
sheet["E3"] = datetime(2026, 9, 24)
sheet["E4"] = True
sheet["E5"] = "Link"
sheet["E5"].hyperlink = "https://example.com"
sheet.merge_cells("H1:I1")
sheet["H1"] = "Merged"
sheet.freeze_panes = "B2"
sheet.auto_filter.ref = "A1:C6"
sheet.row_dimensions[1].height = 30
sheet.column_dimensions["A"].width = 22
fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
rules = [
    CellIsRule(operator="greaterThan", formula=["5"], fill=fill),
    FormulaRule(formula=["B2>0"], font=Font(bold=True)),
    ColorScaleRule(start_type="min", start_color="FF0000", end_type="max", end_color="00FF00"),
    DataBarRule(start_type="min", end_type="max", color="336699"),
    IconSetRule(icon_style="3Arrows", type="percent", values=[0, 33, 67]),
]
for rule in rules:
    sheet.conditional_formatting.add("B2:B6", rule)
for index, kind in enumerate(["list", "whole", "decimal", "date", "time", "textLength", "custom"]):
    rule = DataValidation(type=kind, formula1='"Red,Blue"' if kind == "list" else "B2>0" if kind == "custom" else "1", formula2=None if kind in ["list", "custom"] else "10")
    sheet.add_data_validation(rule)
    rule.add(f"{chr(74 + index)}2:{chr(74 + index)}6")
for index, cls in enumerate([BarChart, LineChart, PieChart]):
    chart = cls()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=6), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=6))
    sheet.add_chart(chart, f"A{10 + index * 16}")
book.create_sheet("Untouched")["A1"] = "Preserve this worksheet"
book.save(path)
with ZipFile(path) as archive:
    parts = {name: archive.read(name) for name in archive.namelist()}
parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(b"</worksheet>", b'<extLst><ext uri="urn:react-forge:fixture"><opaque xmlns="urn:react-forge:fixture">Preserve unsupported rule extensions.</opaque></ext></extLst></worksheet>')
with ZipFile(path, "w", ZIP_DEFLATED) as archive:
    for name, data in parts.items():
        archive.writestr(name, data)
