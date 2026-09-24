"""Test-only fixture producer: python-docx 1.2.0; never a runtime dependency."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from docx import Document
from docx.oxml import OxmlElement

root = Path(__file__).parent
document = Document()
document.add_heading("External title", 1)
document.add_paragraph("External paragraph")
table = document.add_table(rows=2, cols=2)
table.style = "Table Grid"
table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged external header"
table.cell(1, 0).text = "Left"
table.cell(1, 1).text = "Right"
document.sections[0].header.paragraphs[0].text = "External header"
opaque = document.add_paragraph("Protected equation")
equation = OxmlElement("m:oMath")
run = OxmlElement("m:r")
text = OxmlElement("m:t")
text.text = "x = 1"
run.append(text)
equation.append(run)
opaque._p.append(equation)
document.save(root / "external.docx")
with ZipFile(root / "external.docx") as archive:
    parts = {name: archive.read(name) for name in archive.namelist()}
parts["word/unrelated.xml"] = b'<opaque xmlns="urn:react-forge:fixture">Preserve these exact bytes.</opaque>'
with ZipFile(root / "external.docx", "w", ZIP_DEFLATED) as archive:
    for name, data in parts.items():
        archive.writestr(name, data)

# A second independent fixture combines python-docx's document/package writer
# with openpyxl's native chart XML. No React Forge writer produces these inputs.
from docx.opc.constants import RELATIONSHIP_TYPE as RT, CONTENT_TYPE as CT
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.oxml import parse_xml
from lxml import etree

source = root.parents[2] / 'forge-xlsx/tests/fixtures/external.xlsx'
workbook_bytes = source.read_bytes()
charts = Document()
charts.add_paragraph('Before external charts')
charts.add_paragraph('Original list item', style='List Number')
charts.add_picture(str(root.parents[2] / 'forge-pptx/tests/fixtures/sample.png'))
namespace = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
relationships = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
with ZipFile(source) as workbook:
    for i in range(1, 4):
        chart = etree.fromstring(workbook.read(f'xl/charts/chart{i}.xml'))
        embedded = Part(PackURI(f'/word/embeddings/data{i}.xlsx'), CT.SML_SHEET, workbook_bytes, charts.part.package)
        part = Part(PackURI(f'/word/charts/chart{i}.xml'), CT.DML_CHART, b'', charts.part.package)
        data_id = part.relate_to(embedded, RT.PACKAGE)
        external = etree.SubElement(chart, f'{{{namespace}}}externalData', {f'{{{relationships}}}id': data_id})
        etree.SubElement(external, f'{{{namespace}}}autoUpdate', val='0')
        part._blob = etree.tostring(chart, xml_declaration=True, encoding='UTF-8', standalone=True)
        chart_id = charts.part.relate_to(part, RT.CHART)
        paragraph = charts.add_paragraph()
        paragraph._p.append(parse_xml(f'''<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
          xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
          xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:c="{namespace}" xmlns:r="{relationships}">
          <w:drawing><wp:inline><wp:extent cx="4572000" cy="2286000"/><wp:docPr id="{i + 1}" name="External chart {i}" descr="Editable chart"/>
          <a:graphic><a:graphicData uri="{namespace}"><c:chart r:id="{chart_id}"/></a:graphicData></a:graphic></wp:inline></w:drawing></w:r>'''))
charts.add_paragraph('After external charts')
charts.save(root / 'charts.docx')
