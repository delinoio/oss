"""Reproduce the externally authored fixture with python-pptx 1.0.2 and Pillow.
These tools are test-fixture authors only, never Forge runtime dependencies.
"""
from pathlib import Path
from io import BytesIO
from PIL import Image
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.dml.color import RGBColor
from lxml import etree

root=Path(__file__).parent
image=Image.new('RGB',(200,100),(37,99,235))
buf=BytesIO(); image.save(buf,format='PNG'); (root/'sample.png').write_bytes(buf.getvalue())
image=Image.new('RGB',(100,200),(34,197,94)); image.save(root/'replacement.png')
p=Presentation(); p.slide_width=Pt(960); p.slide_height=Pt(540)
s=p.slides.add_slide(p.slide_layouts[5]); s.shapes.title.text='External presentation'
box=s.shapes.add_textbox(Pt(40),Pt(90),Pt(440),Pt(60)); box.text='External text with preserved metadata'
for r in box.text_frame.paragraphs[0].runs:
 r.font.name='Noto Sans KR'; r.font.size=Pt(18); r.font.color.rgb=RGBColor(23,32,51)
 etree.SubElement(r._r.rPr,'{urn:forge-test}extension',{'keep':'yes'})
s.shapes.add_picture(str(root/'sample.png'),Pt(40),Pt(180),width=Pt(180))
t=s.shapes.add_table(3,3,Pt(250),Pt(180),Pt(280),Pt(180)).table
for y in range(3):
 for x in range(3):
  t.cell(y,x).text=f'Cell {y},{x}'
  for r in t.cell(y,x).text_frame.paragraphs[0].runs: r.font.name='Noto Sans KR'; r.font.size=Pt(14)
t.cell(0,0).merge(t.cell(0,2));t.cell(0,0).text='Merged header'
d=CategoryChartData(); d.categories=['A','B','C']; d.add_series('First',[10,20,30]); d.add_series('Second',[15,25,35])
s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Pt(570),Pt(130),Pt(330),Pt(260),d)
group=s.shapes.add_group_shape(); group.shapes.add_shape(MSO_SHAPE.HEXAGON,Pt(40),Pt(410),Pt(120),Pt(70))
s.shapes.add_shape(MSO_SHAPE.STAR_5_POINT,Pt(760),Pt(410),Pt(90),Pt(70))
p.save(root/'external.pptx')
